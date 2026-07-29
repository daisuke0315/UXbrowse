from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import build_medical_physics_draft as base


base.FONT = "Hiragino Sans"
MAIN_PATH = base.OUT_DIR / "Medical_Physics_Manuscript_Draft_Japanese.docx"
TITLE_PATH = base.OUT_DIR / "Medical_Physics_Title_Page_Draft_Japanese.docx"


def ph(doc, text):
    return base.add_placeholder_paragraph(doc, text)


def add_title_page():
    doc = Document()
    base.configure_document(doc, line_numbers=False)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.15
    doc.styles["Normal"].paragraph_format.first_line_indent = Inches(0)
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[style_name].paragraph_format.line_spacing = 1.15
        doc.styles[style_name].paragraph_format.space_before = Pt(7)
        doc.styles[style_name].paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(
        "CT座標拘束型可視光三次元再構成と視点連動CT多断面再構成表示："
        "ファントム精度評価"
    )
    base.set_run_font(r, size=15, bold=True)

    base.add_rich_paragraph(
        doc,
        [{"text": "論文種別：", "bold": True}, {"text": "Research Article"}],
        indent=False,
    )
    base.add_rich_paragraph(
        doc,
        [{"text": "短縮表題：", "bold": True}, {"text": "CT拘束型可視光3Dと視点連動MPR"}],
        indent=False,
    )

    base.add_heading(doc, "著者", 1)
    base.add_rich_paragraph(
        doc,
        [
            {"text": "谷井 大介", "bold": True},
            {"text": "［英文氏名表記・学位を確認］", "placeholder": True},
            {"text": "；［共著者氏名・学位］", "placeholder": True},
        ],
        indent=False,
    )

    base.add_heading(doc, "所属", 1)
    base.add_rich_paragraph(
        doc,
        [
            {"text": "1．札幌孝仁会記念病院 SAFRA"},
            {"text": "［正式な英文所属名を確認］", "placeholder": True},
        ],
        indent=False,
    )
    ph(doc, "共著者全員の所属と所属番号を追記する。")

    base.add_heading(doc, "責任著者", 1)
    ph(doc, "氏名、部署、施設、郵便住所、国、電話番号、メールアドレスを記載する。")

    base.add_heading(doc, "著者貢献", 1)
    base.add_plain(
        doc,
        "D.T.：構想、方法論、ソフトウェア、実験、データ管理、解析、可視化、初稿作成。"
        "［共著者］：妥当性確認、臨床的解釈、研究監督、原稿のレビューおよび編集。",
        indent=False,
    )
    ph(doc, "CRediT分類を確定し、全著者の承認を得る。")

    base.add_heading(doc, "謝辞", 1)
    ph(doc, "著者要件を満たさない協力者を記載する。該当しない場合は「なし」とする。")

    base.add_heading(doc, "研究費", 1)
    base.add_plain(
        doc,
        "本研究は外部研究費の提供を受けず、著者の自己資金および所属施設の既存環境を"
        "用いて実施した。",
        indent=False,
    )
    ph(doc, "投稿前に研究費の記載を確認する。")

    base.add_heading(doc, "利益相反", 1)
    base.add_plain(
        doc,
        "D.T.は、可視光三次元モデルと内部医用画像の統合に関連するPCT国際出願の"
        "発明者および出願人であり、将来の実用化または外部連携に関する金銭的利害を"
        "有する可能性がある。本研究は特定企業から支援を受けずに設計・評価した。"
        "その他の著者には開示すべき利益相反はない。",
        indent=False,
    )
    ph(doc, "PCT国際出願番号を追記し、所属施設および全著者と開示内容を確認する。")

    base.add_heading(doc, "原稿に関する宣言", 1)
    base.add_plain(
        doc,
        "責任著者は、全著者が本原稿を確認し、投稿を承認したこと、ならびに本研究が"
        "他誌で審査中ではないことを確認する。",
        indent=False,
    )

    doc.core_properties.title = "Medical Physics 日本語タイトルページ初稿"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(TITLE_PATH)


def add_abstract(doc):
    base.add_heading(doc, "要旨", 1)
    base.add_abstract_label(
        doc,
        "背景",
        "Computed tomography（CT）は内部解剖を物理的な患者座標系で表現するが、"
        "患者固有の可視表面を直接保持しない。可視光三次元再構成は体表の外観を保持する"
        "一方、通常は絶対スケールおよびDigital Imaging and Communications in Medicine"
        "（DICOM）座標との対応を持たない。",
    )
    base.add_abstract_label(
        doc,
        "目的",
        "CT由来のGround Control Point（GCP）をバンドル調整へ導入してCT座標系に拘束"
        "された可視光三次元モデルを構築し、観察者の仮想視点に連動する多断面再構成"
        "（MPR）表示を統合したシステムを開発して技術的精度を評価する。",
    )
    base.add_abstract_label(
        doc,
        "方法",
        "CT画像および可視光画像の両方で視認可能な基準マーカーをファントム周囲に配置した。"
        "DICOM物理空間で計測したマーカー座標をGCP拘束としてSfM/MVS再構成へ導入した。"
        "仮想カメラの位置、注視点および上方向ベクトルから断面を定義し、CT値を物理座標上で"
        "再標本化して視点連動MPRを生成した。約40 cmに分布する4点を3回の再構成で評価し、"
        "形状一致性12点を得た。さらに、約30 cmに分布する非GCPの4点を用い、GCP拘束あり・"
        "なしの各3モデル、各条件12点でEnd-to-End Target Registration Error（TRE）を評価した。",
    )
    base.add_abstract_label(
        doc,
        "結果",
        "位置合わせ後の三次元幾何点誤差は0.27±0.12 mmであった。End-to-End TREはGCP拘束"
        "ありで0.45±0.26 mm、拘束なしで1.47±0.95 mmであり、平均TREは1.02 mm（69.4%）"
        "低減した。",
    )
    base.add_abstract_label(
        doc,
        "結論",
        "CT由来GCPを可視光再構成へ導入することで、静的ファントム条件下において可視光"
        "三次元モデルと視点連動CT MPRをサブミリメートル水準で統合できた。ただし、"
        "Procrustes解析およびICP後の幾何誤差は未補正の絶対座標精度ではなく形状一致性を"
        "表す。再現性、不確かさ、運動感度および人体での性能について追加検証が必要である。",
    )
    base.add_rich_paragraph(
        doc,
        [
            {"text": "キーワード：", "bold": True},
            {"text": "CT、可視光三次元再構成、Ground Control Point、Bundle Adjustment、"
                     "画像位置合わせ、MPR"},
        ],
        indent=False,
    )


def add_main_manuscript():
    doc = Document()
    base.configure_document(doc, line_numbers=True)
    doc.core_properties.title = "Medical Physics 日本語原稿初稿"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = "黄色の未確定項目を含む作業用初稿"

    hp = doc.sections[0].header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = hp.add_run("作業用初稿―投稿不可")
    base.set_run_font(hr, size=9, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Inches(0)
    title.paragraph_format.space_after = Pt(10)
    r = title.add_run(
        "CT座標拘束型可視光三次元再構成と視点連動CT多断面再構成表示："
        "ファントム精度評価"
    )
    base.set_run_font(r, size=15, bold=True)

    add_abstract(doc)

    base.add_heading(doc, "1．緒言", 1)
    base.add_plain(
        doc,
        "CTは、内部の解剖学的構造を物理的な患者座標系とともに高精度に表現し、診断、"
        "治療計画および治療効果判定に広く用いられている。一方、通常の断面表示や"
        "ボリュームレンダリングは、皮膚色、体表ランドマーク、表面テクスチャなどの"
        "患者固有の可視光外観を保持しない。このため、内部の標的と体表上の位置との関係は、"
        "非専門家にとって理解しにくい場合がある。¹˒²",
    )
    base.add_plain(
        doc,
        "Structure from Motion（SfM）、Multiview Stereo（MVS）、Neural Radiance Field、"
        "3D Gaussian Splattingなどは、可視光画像から写実的な三次元外観を再構成できる。³˒⁴ "
        "しかし、画像対応のみから得られる再構成は一般に相似変換の自由度を持ち、CTデータと"
        "同一のスケール、方向、原点または患者座標を本質的には共有しない。完成後のモデルを"
        "剛体または相似変換でCTへ登録する方法では、再構成中のスケールドリフト、カメラ姿勢"
        "誤差または局所変形を直接拘束できない。⁴˒⁵",
    )
    base.add_plain(
        doc,
        "GCPは既知の三次元座標としてバンドル調整へ導入でき、その数と空間配置が全体精度に"
        "影響する。⁵ 医用画像分野でも、光学表面モデルとCTまたはCone-beam CTとの位置合わせが"
        "検討されている。⁶ 本研究では、CT由来GCPを再構成段階へ導入し、可視光モデルをDICOM"
        "物理空間に生成するとともに、観察視点から任意CT断面を生成するシステムを開発した。"
        "静的ファントムを用いて、位置合わせ後の形状一致性とEnd-to-End TREを評価し、"
        "GCP拘束なしの再構成と比較した。",
    )

    base.add_heading(doc, "2．材料と方法", 1)
    base.add_heading(doc, "2.1　システム構成", 2)
    base.add_plain(
        doc,
        "システムは、CT・可視光画像取得、GCP拘束付き可視光三次元再構成、DICOM物理座標に"
        "おける任意断面MPR生成、可視光モデルとMPRの統合描画、の4段階で構成した。主要な"
        "構成要素を表Iに、処理フローを図1に示す。CT由来マーカー座標と可視光再構成は共通の"
        "右手系物理座標で表現し、座標軸、変換行列および単位変換を処理過程で保持した。",
    )
    ph(doc, "座標系の右手・左手、患者方向、変換順序、軸入れ替えおよび符号反転の有無を確認する。")

    base.add_heading(doc, "2.2　ファントムおよび基準マーカー", 2)
    base.add_plain(
        doc,
        "静的ファントムの周囲に、CT画像と可視光画像の両方で位置を同定できるマーカーを"
        "配置した。一部を再構成用GCPとし、別のランドマークを精度評価用に確保した。"
        "各マーカー中心をCT物理座標と可視光モデル上で同定し、制御点と評価点が再構成範囲に"
        "広く分布するよう配置した。",
    )
    ph(doc, "ファントム名、メーカー、材質、寸法、マーカー材質・直径、総数、GCP数、評価点数および点間距離図を追記する。")
    ph(doc, "形状評価の4点がGCPと独立しているか確認する。GCPそのものであれば評価名称と解釈を修正する。")

    base.add_heading(doc, "2.3　CT撮像およびランドマーク計測", 2)
    base.add_plain(
        doc,
        "臨床用CT装置でファントムを撮像し、Image Position (Patient)、Image Orientation "
        "(Patient)、Pixel Spacingおよびスライス位置を含むDICOM情報を保持した。マーカー中心の"
        "物理座標をCTボリュームから取得した。MPR生成時には、任意物理座標のCT値を三線形補間"
        "により求めた。",
    )
    ph(doc, "CTメーカー・機種、撮像モード、管電圧、管電流またはmAs、ピッチ、コリメーション、再構成スライス厚・間隔、FOV、マトリクス、再構成関数、ボクセル寸法を追記する。")
    ph(doc, "マーカー中心の決定方法、観察者経験、反復計測回数、観察者内・観察者間の位置同定不確かさを記述する。")

    base.add_heading(doc, "2.4　可視光撮像および三次元再構成", 2)
    base.add_plain(
        doc,
        "固定したファントムの周囲でiPhone 15 Pro（Apple Inc.）を移動させ、可視光動画を"
        "取得した。FFmpegでフレームを抽出し、SfM/MVS処理によりカメラ内部パラメータ、"
        "カメラ姿勢および三次元構造を推定した。GCPの画像観測とDICOM物理座標を関連付け、"
        "最終バンドル調整前に導入した。SfM/MVSにはRealityCapture、3D Gaussian表現を生成する"
        "場合にはPostShotを使用した。",
    )
    ph(doc, "動画解像度、フレームレート、露出・焦点、撮影距離、軌道、照明、抽出フレーム数、選択基準、全ソフトウェアのバージョンと再構成設定を追記する。")
    ph(doc, "精度計測をメッシュ、点群、Gaussian表現のどれで行ったか、およびマーカー中心の選択方法を記述する。")

    base.add_heading(doc, "2.5　CT座標拘束付きバンドル調整", 2)
    base.add_plain(
        doc,
        "通常のバンドル調整は、観測画像座標と推定三次元点の投影位置との差を最小化する。"
        "本研究では、再構成されたGCP位置と既知CT座標との差に対する項を追加した。",
    )
    eq = doc.add_paragraph()
    eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    eq.paragraph_format.first_line_indent = Inches(0)
    er = eq.add_run("min Σᵢⱼ ‖π(Kᵢ, Rᵢ, tᵢ, Xⱼ) − xᵢⱼ‖² + λ Σₖ ‖Xgₖ − Gₖ‖²")
    base.set_run_font(er, name="Cambria Math", size=11)
    base.add_plain(
        doc,
        "ここで、xᵢⱼは画像iにおける特徴点jの観測位置、Xⱼは推定三次元点、Kᵢ、Rᵢ、tᵢは"
        "カメラiの内部行列、回転、並進、GₖはGCP kの既知CT座標、Xgₖはその再構成座標、"
        "πは投影、λは座標拘束の相対重みである。GCP項により、再構成時の並進、回転および"
        "スケールを拘束した。同一可視光画像からGCP項を用いない非拘束モデルも生成した。",
    )
    ph(doc, "GCP重みまたは設定座標精度、残差閾値、外れ値除外、最終再投影誤差、GCP残差、非拘束条件の位置合わせ方法を報告する。")

    base.add_heading(doc, "2.6　視点連動MPR生成", 2)
    base.add_plain(
        doc,
        "仮想カメラ位置をeye、注視点をtarget、画面上方向をupとした。断面法線nと面内直交軸"
        "u、vを次式で求めた。",
    )
    for text in (
        "n = normalize(target − eye)",
        "u = normalize(up × n)",
        "v = n × u",
        "X(i, j) = center + (i − cᵢ)Δu·u + (j − cⱼ)Δv·v",
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
        rr = p.add_run(text)
        base.set_run_font(rr, name="Cambria Math", size=11)
    base.add_plain(
        doc,
        "cᵢ、cⱼは画像中心、Δu、Δvは面内画素間隔である。DICOM幾何情報を用いて物理座標を"
        "連続ボクセル座標へ変換し、三線形補間でCT値を取得した。eye、target、upを表示側から"
        "WebSocketで送信し、生成したMPR画像を対応断面へテクスチャとして配置した。",
    )
    ph(doc, "MPRマトリクス、画素間隔、FOV、補間実装、Window/Level、更新頻度、遅延計測、CT範囲外の処理を追記する。")

    base.add_heading(doc, "2.7　統合描画", 2)
    base.add_plain(
        doc,
        "可視光三次元モデルとCTボリュームを同一物理座標系で保持した。観察視点の変更に"
        "応じてMPR断面を更新し、外観モデルのナビゲーションと内部CTのナビゲーションを"
        "連動させた。DICOM処理および数値演算にはPython、SimpleITK、NumPy、SciPy、通信には"
        "FastAPIとWebSocket、描画とUIにはUnityを用いた。",
    )
    ph(doc, "各ソフトウェア・エンジンのバージョン、ワークステーション、GPU、OS、Unityへのモデル取り込み時の変換を追記する。")

    base.add_heading(doc, "2.8　精度評価", 2)
    base.add_heading(doc, "2.8.1　位置合わせ後の形状一致性", 3)
    base.add_plain(
        doc,
        "約40 cmに分布する4対応点を用い、3回の再構成から12点を評価した。重心、スケール、"
        "特異値分解による回転および並進を補正するProcrustes解析後、ICPで最適化し、対応点間"
        "ユークリッド距離を算出した。全体の座標変換を除去しているため、本指標は位置合わせ後"
        "の形状一致性であり、未補正のCT絶対座標精度ではない。",
    )
    ph(doc, "Procrustes/ICP前の絶対座標誤差を追加するか、形状一致性に限定する明確な理由を記述する。")

    base.add_heading(doc, "2.8.2　End-to-End統合表示精度", 3)
    base.add_plain(
        doc,
        "GCPに使用していない約30 cmに分布する4ランドマークを独立標的とした。可視光モデル"
        "上のマーカー中心と統合MPR上の対応位置との三次元ユークリッド距離をTREとした。"
        "GCP拘束あり・なしの各条件で3モデルを評価し、各条件12点を得た。この計測には、"
        "可視光撮像、再構成、座標転送、MPR生成および統合描画の全工程が含まれた。",
    )
    ph(doc, "計測ツール、観察者、盲検化、反復点選択、ボクセルから表示座標への変換、両条件の事後位置合わせ、対応データか否かを記述する。")

    base.add_heading(doc, "2.9　統計解析", 2)
    base.add_plain(
        doc,
        "本研究は初期技術検証であるため記述統計を用い、誤差を平均±標準偏差で示した。"
        "平均TRE低減率は（非拘束平均TRE−拘束平均TRE）/非拘束平均TRE×100%で算出した。"
        "再構成モデルが3つであり、点データの独立性と対応構造を原データで確認する必要が"
        "あるため、推測統計は実施しなかった。解析にはPythonを用いた。",
    )
    ph(doc, "個別値、中央値、範囲、最大誤差、95%信頼区間、再構成内で反復する点を考慮した解析、Python・パッケージのバージョンを追記する。")

    base.add_heading(doc, "2.10　倫理的配慮", 2)
    base.add_plain(
        doc,
        "本研究は静的人体模擬ファントムおよび研究用画像データのみを用い、人を対象とする"
        "研究、患者画像、診療情報、生体試料または介入を含まなかった。",
    )
    ph(doc, "所属施設で倫理審査不要または免除の正式確認が必要か確認し、必要なら番号を記載する。")

    base.add_heading(doc, "2.11　生成AI使用の開示", 2)
    base.add_plain(
        doc,
        "OpenAI Codexを、原稿の英語化、文章構成および書式調整の補助に使用した。実験データの"
        "生成、報告した解析または科学的判断には使用していない。著者が内容を確認し、"
        "原稿全体に責任を負う。",
    )

    base.add_heading(doc, "3．結果", 1)
    base.add_heading(doc, "3.1　再構成およびシステム動作", 2)
    base.add_plain(
        doc,
        "GCP拘束付きワークフローにより、CT物理座標系上の可視光三次元モデルを生成し、"
        "視点に連動するMPR画像を対応断面へ配置できた。代表表示を図2に示す。",
    )
    ph(doc, "試行数、成功数、失敗基準、工程別処理時間、描画フレームレート、MPR応答遅延を報告する。")

    base.add_heading(doc, "3.2　位置合わせ後の形状一致性", 2)
    base.add_plain(
        doc,
        "3回の再構成、各4点の計12点における位置合わせ後の三次元幾何点誤差は"
        "0.27±0.12 mmであった。評価点は約40 cmの範囲に分布した。点配置と解析手順を"
        "図3に示す。",
    )
    ph(doc, "12点の個別値、中央値、範囲、最大値、95%信頼区間を追加し、位置合わせ前の絶対座標誤差を別途報告する。")

    base.add_heading(doc, "3.3　End-to-End統合表示精度", 2)
    base.add_plain(
        doc,
        "End-to-End TREは、GCP拘束ありで0.45±0.26 mm、拘束なしで1.47±0.95 mmであった"
        "（各条件12点）。平均差は1.02 mmで、相対低減率は69.4%であった。結果を表IIに示す。",
    )
    ph(doc, "対応する原データを追加し、散布図または箱ひげ図を完成させる。最大誤差と信頼区間を報告し、統計検定がない限り「有意」という表現を用いない。")

    base.add_heading(doc, "4．考察", 1)
    base.add_plain(
        doc,
        "本研究では、CT由来の物理座標を可視光バンドル調整へ導入し、可視光三次元モデルと"
        "視点連動CT MPRを統合した。静的ファントム条件におけるEnd-to-End TREは、GCP拘束"
        "ありで0.45 mm、拘束なしで1.47 mmであった。この結果は、再構成段階で座標情報を"
        "導入することが、可視光外観モデルとCT表示の下流整合性を改善し得ることを示す。",
    )
    base.add_plain(
        doc,
        "画像対応のみの再構成は、再投影誤差が小さくても物理スケールや患者座標との一致を"
        "保証しない。GCP観測は、再構成中の全体相似変換自由度を拘束し、カメラおよび点推定を"
        "安定化する可能性がある。ただし、再構成回数が少ないため、効果量と変動性を精密には"
        "推定できない。",
    )
    base.add_plain(
        doc,
        "2つの精度指標は分けて解釈する必要がある。0.27 mmはスケール補正を含むProcrustes"
        "解析とICP後の値であり、全体座標変換を除いた残余形状誤差を表す。一方、End-to-End "
        "TREは非GCPランドマークにおける統合表示全体の精度を意図した。CT座標拘束の絶対精度"
        "を確立するには、未補正座標誤差、GCP残差、独立標的誤差の追加報告が必要である。",
    )
    base.add_plain(
        doc,
        "視点連動MPRでは、仮想カメラ姿勢がMPR基底を直接定義するため、観察中の外観部位と"
        "内部CT断面を幾何学的に連動できる。画像誘導計画、内外位置関係の説明または教育への"
        "応用が考えられるが、本研究は診断性能、処置誘導、患者理解または臨床安全性を評価して"
        "いない。用途ごとに必要精度と検証項目を別途定める必要がある。",
    )

    base.add_heading(doc, "4.1　限界", 2)
    limitations = [
        "評価は3モデル、各4ランドマークに限定された。",
        "同一再構成内の複数ランドマークは統計的に独立とは限らない。",
        "静的ファントムであり、呼吸、体動、軟部組織変形、撮像間姿勢差を含まない。",
        "GCP数、配置、材質、視認性および位置同定不確かさを系統的に変化させていない。",
        "形状評価では全体スケールと姿勢誤差を除去しており、絶対座標精度を示さない。",
        "最大誤差、信頼区間、観察者差、装置差、日間再現性を評価していない。",
        "メッシュ、点群、視点依存Gaussian表現の違いがランドマーク計測へ与える影響を評価していない。",
        "処理時間、失敗条件、品質管理閾値および臨床ワークフローを正式に評価していない。",
    ]
    for item in limitations:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.5
        rr = p.add_run(item)
        base.set_run_font(rr)

    base.add_heading(doc, "4.2　今後の検証", 2)
    base.add_plain(
        doc,
        "次の技術研究では、独立した再構成回数を増やし、誤差予算を作成するとともに、GCP数・"
        "配置、マーカー位置同定誤差、照明、撮影軌道、再構成条件を評価する必要がある。"
        "拘束付き再構成を、事前に規定した剛体・相似位置合わせ法と比較する。その後、人を"
        "対象として再現性、空間理解、回答時間および臨床的操作性を別々に検証する。",
    )

    base.add_heading(doc, "5．結論", 1)
    base.add_plain(
        doc,
        "可視光バンドル調整へCT由来GCPを導入することで、可視光三次元モデルと視点連動CT "
        "MPRを統合できた。静的ファントムにおける平均End-to-End TREは、GCP拘束ありで"
        "0.45 mm、拘束なしで1.47 mmであった。本結果は技術的成立性を示すが、臨床性能を"
        "示すものではない。絶対座標精度、再現性、不確かさ、運動感度および人体での性能を"
        "追加検証する必要がある。",
    )

    base.add_heading(doc, "データ利用可能性", 1)
    ph(doc, "個別ランドマーク誤差および撮像メタデータを含む最小データセットの公開先、またはアクセス制限の理由を記載する。")

    base.add_heading(doc, "コード利用可能性", 1)
    ph(doc, "解析・MPR生成コードの公開範囲、リポジトリ、バージョンまたはDOI、商用ソフトウェア・特許による制限を記載する。")

    base.add_heading(doc, "参考文献", 1)
    references = [
        "Dappa E, Higashigaito K, Fornaro J, et al. Cinematic rendering—an alternative to volume rendering for 3D computed tomography imaging. Insights Imaging. 2016;7(6):849-856. doi:10.1007/s13244-016-0518-1.",
        "Pose-Díez-de-la-Lastra A, Ungi T, Morton D, et al. Real-time integration between Microsoft HoloLens 2 and 3D Slicer with demonstration in pedicle screw placement planning. Int J Comput Assist Radiol Surg. 2023;18:2023-2032. doi:10.1007/s11548-023-02977-0.",
        "Kerbl B, Kopanas G, Leimkühler T, Drettakis G. 3D Gaussian splatting for real-time radiance field rendering. ACM Trans Graph. 2023;42(4):139. doi:10.1145/3592433.",
        "Triggs B, McLauchlan PF, Hartley RI, Fitzgibbon AW. Bundle adjustment—a modern synthesis. In: Triggs B, Zisserman A, Szeliski R, eds. Vision Algorithms: Theory and Practice. Springer; 2000:298-372.",
        "Sanz-Ablanedo E, Chandler JH, Rodríguez-Pérez JR, Ordóñez C. Accuracy of unmanned aerial vehicle and SfM photogrammetry survey as a function of the number and location of ground control points used. Remote Sens. 2018;10(10):1606. doi:10.3390/rs10101606.",
        "Lim SW, et al. Registration accuracy between intraoral-scanned and cone-beam computed tomography-scanned crowns in various registration methods. Am J Orthod Dentofacial Orthop. 2020;157(3):348-356. doi:10.1016/j.ajodo.2019.04.031.",
        "Phelps EE, Wellings R, Griffiths F, et al. Do medical images aid understanding and recall of medical information? Patient Educ Couns. 2017;100(6):1120-1127.",
        "Stewart JA, et al. Visual teaching aids improve patient understanding and reduce anxiety prior to a colectomy. Am J Surg. 2021;222(4):780-785.",
        "Julious SA. Sample size of 12 per group rule of thumb for a pilot study. Pharm Stat. 2005;4(4):287-291. doi:10.1002/pst.185.",
    ]
    for ref in references:
        base.add_reference(doc, ref)
    ph(doc, "全参考文献をPubMed/Crossrefで確認し、AMA様式に必要な著者情報を補完する。特に文献6と8を確認する。")

    base.add_heading(doc, "表", 1)
    base.add_plain(doc, "表I．システム構成と機能。", style="Caption", indent=False)
    table1 = doc.add_table(rows=1, cols=3)
    table1.style = "Table Grid"
    widths1 = [1900, 3000, 4460]
    for idx, text in enumerate(("機能", "機器・ソフトウェア", "役割")):
        base.fill_table_cell(table1.rows[0].cells[idx], text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        base.set_cell_shading(table1.rows[0].cells[idx], base.TABLE_HEADER_FILL)
    base.mark_table_header(table1.rows[0])
    rows1 = [
        ("可視光撮像", "iPhone 15 Pro", "多視点可視光動画の取得"),
        ("フレーム抽出", "FFmpeg［バージョン］", "静止画フレームの抽出と選択"),
        ("三次元再構成", "RealityCapture［バージョン］", "SfM/MVSおよびGCP拘束付きバンドル調整"),
        ("三次元外観モデル", "PostShot［バージョン］", "3D Gaussian表現の生成と描画"),
        ("DICOM処理", "Python、SimpleITK、NumPy、SciPy［各バージョン］", "物理座標変換、補間、MPR生成"),
        ("通信", "FastAPI、WebSocket［各バージョン］", "視点情報とMPR画像の双方向通信"),
        ("描画・UI", "Unity［バージョン］", "可視光モデル表示、視点制御、MPR重畳"),
    ]
    for row_data in rows1:
        cells = table1.add_row().cells
        for idx, value in enumerate(row_data):
            base.fill_table_cell(cells[idx], value, placeholder="［" in value)
    base.set_table_geometry(table1, widths1)

    base.add_plain(doc, "表II．技術的精度評価結果。", style="Caption", indent=False)
    table2 = doc.add_table(rows=1, cols=5)
    table2.style = "Table Grid"
    widths2 = [2200, 2200, 1000, 1500, 2460]
    for idx, text in enumerate(("評価指標", "条件", "モデル数", "評価点数", "誤差（平均±SD）")):
        base.fill_table_cell(
            table2.rows[0].cells[idx],
            text,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            size=9,
        )
        base.set_cell_shading(table2.rows[0].cells[idx], base.TABLE_HEADER_FILL)
    base.mark_table_header(table2.rows[0])
    rows2 = [
        ("位置合わせ後の幾何点誤差", "GCP拘束あり", "3", "12", "0.27±0.12 mm"),
        ("End-to-End TRE", "GCP拘束あり", "3", "12", "0.45±0.26 mm"),
        ("End-to-End TRE", "GCP拘束なし", "3", "12", "1.47±0.95 mm"),
    ]
    for row_data in rows2:
        cells = table2.add_row().cells
        for idx, value in enumerate(row_data):
            align = WD_ALIGN_PARAGRAPH.LEFT if idx < 2 else WD_ALIGN_PARAGRAPH.CENTER
            base.fill_table_cell(cells[idx], value, align=align)
    base.set_table_geometry(table2, widths2)

    base.add_heading(doc, "図説明", 1)
    legends = [
        ("図1．", "CT・可視光画像取得、CT由来GCP同定、拘束付き可視光再構成、DICOM座標MPR生成、統合描画までの処理フロー。"),
        ("図2．", "統合表示の代表例。可視光三次元モデルと視点連動CT MPR断面を同一物理座標系で表示する。外観、MPR断面、座標軸を示すパネルを追加する。"),
        ("図3．", "幾何精度評価。ファントム、約40 cmの点配置、対応点定義、未補正絶対誤差と位置合わせ後形状誤差の違いを示す。"),
        ("図4．", "End-to-End TRE評価。約30 cmに分布する4つの非GCP標的と、可視光ランドマークからMPR対応位置までの計測を示す。"),
        ("図5．", "GCP拘束あり・なしにおけるEnd-to-End TRE分布。全個別値とモデル／再構成のまとまりを表示し、箱、ひげ、信頼区間の定義を最終説明へ記載する。"),
    ]
    for label, legend in legends:
        p = doc.add_paragraph(style="Caption")
        r1 = p.add_run(label)
        base.set_run_font(r1, size=10, bold=True)
        r2 = p.add_run(legend)
        base.set_run_font(r2, size=10, bold=False)

    base.add_heading(doc, "作業用完成チェック（投稿前に削除）", 1)
    checklist = [
        "黄色の未確定項目を検証済み情報で置き換える。",
        "位置合わせ前の絶対座標精度と点ごとの原データを追加する。",
        "観測値の独立性・対応構造を確認し、統計解析を再実施する。",
        "図1～5を投稿品質で完成させる。",
        "参考文献を確認し、表面―CT位置合わせと不確かさに関する文献を追加する。",
        "科学的レビューおよび共著者確認後にチェックリストと作業用ヘッダーを削除する。",
    ]
    for item in checklist:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.line_spacing = 1.5
        rr = p.add_run(item)
        base.set_run_font(rr)

    doc.save(MAIN_PATH)


if __name__ == "__main__":
    add_title_page()
    add_main_manuscript()
    print(MAIN_PATH)
    print(TITLE_PATH)
